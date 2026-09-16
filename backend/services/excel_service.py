import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from backend.config import EXCEL_OUTPUT_DIR
from backend.models import AttendanceSession, User


def generate_attendance_excel(session_id: int, db_session) -> str:
    session = db_session.query(AttendanceSession).filter(
        AttendanceSession.id == session_id
    ).first()

    if not session:
        raise ValueError(f"Attendance session {session_id} not found.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Sheet"
    ws.views.sheetView[0].showGridLines = True

    # Palette
    NAVY_HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    SUB_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="1E293B")
    LABEL_FONT = Font(name="Calibri", size=11, bold=True, color="475569")
    VAL_FONT = Font(name="Calibri", size=11, bold=False, color="0F172A")
    TABLE_HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    TABLE_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    PRESENT_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    PRESENT_FONT = Font(name="Calibri", size=11, bold=True, color="166534")
    ABSENT_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    ABSENT_FONT = Font(name="Calibri", size=11, bold=True, color="991B1B")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # 1. Main Title Header Banner
    ws.merge_cells('A1:I1')
    title_cell = ws['A1']
    title_cell.value = "INSTITUTION AUTOMATED ATTENDANCE REPORT"
    title_cell.font = HEADER_FONT
    title_cell.fill = NAVY_HEADER_FILL
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # 2. Institutional Metadata Section
    dept_name = session.class_room.department.name if session.class_room and session.class_room.department else "AI & DS"
    class_name = session.class_room.name if session.class_room else "N/A"
    subject_name = f"{session.subject.name} ({session.subject.code})" if session.subject else "N/A"
    date_str = session.session_date
    total = session.total_students or 0
    present = session.present_count or 0
    absent = session.absent_count or 0
    pct = round((present / total * 100), 2) if total > 0 else 0.0

    confirmed_by_user = db_session.query(User).filter(User.id == session.confirmed_by_user_id).first()
    confirmed_by_name = confirmed_by_user.full_name if confirmed_by_user else "System / Staff"

    metadata_rows = [
        ("Department:", dept_name, "Date:", date_str),
        ("Class:", class_name, "Total Students:", total),
        ("Subject:", subject_name, "Present:", present),
        ("Staff Confirmed:", confirmed_by_name, "Absent:", absent),
        ("", "", "Attendance Percentage:", f"{pct}%")
    ]

    r_idx = 3
    for m in metadata_rows:
        ws.cell(row=r_idx, column=1, value=m[0]).font = LABEL_FONT
        ws.cell(row=r_idx, column=2, value=m[1]).font = VAL_FONT
        ws.cell(row=r_idx, column=4, value=m[2]).font = LABEL_FONT
        ws.cell(row=r_idx, column=5, value=m[3]).font = VAL_FONT
        ws.row_dimensions[r_idx].height = 20
        r_idx += 1

    # Blank row
    r_idx += 1

    # 3. Student Roster Table Headers
    headers = ["S.No", "Student ID", "Student Name", "Class", "Date", "Subject", "Status", "Recognition Confidence", "Confirmed By"]
    ws.row_dimensions[r_idx].height = 28

    for c_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=r_idx, column=c_idx, value=h_text)
        cell.font = TABLE_HEADER_FONT
        cell.fill = TABLE_HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    r_idx += 1

    # 4. Table Data Rows
    records = session.records
    records_sorted = sorted(records, key=lambda x: x.student.student_id if x.student else "")

    for s_no, rec in enumerate(records_sorted, start=1):
        ws.row_dimensions[r_idx].height = 22
        st_id = rec.student.student_id if rec.student else "N/A"
        st_name = rec.student.name if rec.student else "N/A"
        st_status = rec.status
        st_conf = f"{rec.confidence}%"
        conf_user = db_session.query(User).filter(User.id == rec.confirmed_by_user_id).first()
        conf_by = conf_user.full_name if conf_user else "Auto Recognition"

        row_data = [s_no, st_id, st_name, class_name, date_str, session.subject.name if session.subject else "N/A", st_status, st_conf, conf_by]

        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = VAL_FONT
            cell.border = thin_border
            if c_idx in [1, 2, 4, 5, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Status highlight styling
            if c_idx == 7:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if st_status == "PRESENT":
                    cell.fill = PRESENT_FILL
                    cell.font = PRESENT_FONT
                else:
                    cell.fill = ABSENT_FILL
                    cell.font = ABSENT_FONT

        r_idx += 1

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if cell.row == 1:
                continue
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    file_name = f"Attendance_{class_name.replace(' ', '_')}_{session.session_date}_{session_id}.xlsx"
    output_path = EXCEL_OUTPUT_DIR / file_name
    wb.save(str(output_path))

    return file_name
